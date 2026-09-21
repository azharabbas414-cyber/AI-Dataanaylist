import io
from datetime import datetime

import numpy as np
import pandas as pd

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)
from docx import Document
from docx.shared import Inches, Pt


# ============================================================
# DATA SUMMARY
# ============================================================

def build_report_data(df):
    """Build a general-purpose analytics summary."""

    rows = len(df)
    columns = len(df.columns)

    total_cells = rows * columns
    missing_cells = int(df.isna().sum().sum())
    duplicate_rows = int(df.duplicated().sum())

    completeness = (
        ((total_cells - missing_cells) / total_cells) * 100
        if total_cells
        else 100
    )

    numeric_columns = list(
        df.select_dtypes(include=np.number).columns
    )

    categorical_columns = list(
        df.select_dtypes(
            include=["object", "category", "bool"]
        ).columns
    )

    datetime_columns = list(
        df.select_dtypes(
            include=["datetime", "datetimetz"]
        ).columns
    )

    return {
        "rows": rows,
        "columns": columns,
        "missing_cells": missing_cells,
        "duplicate_rows": duplicate_rows,
        "completeness": completeness,
        "numeric_columns": numeric_columns,
        "categorical_columns": categorical_columns,
        "datetime_columns": datetime_columns,
    }


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================

def generate_executive_summary(df, report_data=None):

    if report_data is None:
        report_data = build_report_data(df)

    summary = []

    summary.append(
        f"The dataset contains {report_data['rows']:,} rows "
        f"and {report_data['columns']:,} columns."
    )

    summary.append(
        f"Overall data completeness is "
        f"{report_data['completeness']:.1f}%."
    )

    if report_data["missing_cells"]:

        summary.append(
            f"There are {report_data['missing_cells']:,} "
            f"missing cells requiring review."
        )

    else:

        summary.append(
            "No missing cells were detected."
        )

    if report_data["duplicate_rows"]:

        summary.append(
            f"{report_data['duplicate_rows']:,} duplicate rows "
            f"were identified."
        )

    else:

        summary.append(
            "No duplicate rows were identified."
        )

    if report_data["numeric_columns"]:

        summary.append(
            f"The dataset contains "
            f"{len(report_data['numeric_columns'])} numeric fields."
        )

    if report_data["categorical_columns"]:

        summary.append(
            f"The dataset contains "
            f"{len(report_data['categorical_columns'])} categorical fields."
        )

    return summary


# ============================================================
# FINDINGS
# ============================================================

def generate_report_findings(df):

    findings = []

    missing = df.isna().sum()

    missing = missing[missing > 0]

    if not missing.empty:

        top_missing = missing.sort_values(
            ascending=False
        ).head(5)

        for column, count in top_missing.items():

            percentage = (
                count / len(df) * 100
                if len(df)
                else 0
            )

            findings.append(
                f"{column} contains {count:,} missing values "
                f"({percentage:.1f}% of rows)."
            )

    duplicates = int(df.duplicated().sum())

    if duplicates:

        findings.append(
            f"The dataset contains {duplicates:,} duplicate rows."
        )

    for column in df.select_dtypes(
        include=np.number
    ).columns:

        values = pd.to_numeric(
            df[column],
            errors="coerce"
        ).dropna()

        if len(values) < 5:
            continue

        q1 = values.quantile(0.25)
        q3 = values.quantile(0.75)

        iqr = q3 - q1

        if iqr == 0:
            continue

        lower = q1 - 1.5 * iqr
        upper = q3 + 1.5 * iqr

        outlier_count = int(
            ((values < lower) | (values > upper)).sum()
        )

        if outlier_count:

            findings.append(
                f"{column} contains approximately "
                f"{outlier_count:,} potential statistical outliers."
            )

    if not findings:

        findings.append(
            "No major automatic data-quality findings were identified."
        )

    return findings


# ============================================================
# NUMERIC STATISTICS
# ============================================================

def numeric_statistics(df):

    numeric_df = df.select_dtypes(
        include=np.number
    )

    if numeric_df.empty:
        return pd.DataFrame()

    return numeric_df.describe().T.reset_index().rename(
        columns={"index": "Column"}
    )


# ============================================================
# CATEGORY SUMMARY
# ============================================================

def categorical_statistics(df):

    results = []

    for column in df.select_dtypes(
        include=["object", "category", "bool"]
    ).columns:

        results.append(
            {
                "Column": column,
                "Unique Values": int(
                    df[column].nunique(dropna=True)
                ),
                "Missing": int(
                    df[column].isna().sum()
                ),
                "Top Value": (
                    df[column]
                    .value_counts(dropna=True)
                    .index[0]
                    if not df[column]
                    .value_counts(dropna=True)
                    .empty
                    else ""
                ),
            }
        )

    return pd.DataFrame(results)


# ============================================================
# CORRELATION SUMMARY
# ============================================================

def correlation_summary(df):

    numeric_df = df.select_dtypes(
        include=np.number
    )

    if numeric_df.shape[1] < 2:
        return pd.DataFrame()

    corr = numeric_df.corr()

    pairs = []

    columns = list(corr.columns)

    for i in range(len(columns)):

        for j in range(i + 1, len(columns)):

            value = corr.iloc[i, j]

            if pd.notna(value):

                pairs.append(
                    {
                        "Field 1": columns[i],
                        "Field 2": columns[j],
                        "Correlation": round(float(value), 3),
                    }
                )

    result = pd.DataFrame(pairs)

    if not result.empty:

        result["Absolute"] = result["Correlation"].abs()

        result = result.sort_values(
            "Absolute",
            ascending=False
        ).drop(
            columns=["Absolute"]
        )

    return result


# ============================================================
# PDF HELPERS
# ============================================================

def _pdf_table(data, widths=None):

    table = Table(
        data,
        colWidths=widths,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#1e293b"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor("#cbd5e1"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#f8fafc"),
                    ],
                ),
            ]
        )
    )

    return table


def _safe_text(value):

    if pd.isna(value):
        return ""

    return str(value)


# ============================================================
# PDF REPORT
# ============================================================

def generate_pdf_report(
    df,
    dataset_name="InsightAI Dataset",
):

    report_data = build_report_data(df)

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontSize=24,
        leading=28,
        alignment=TA_CENTER,
        spaceAfter=18,
    )

    heading_style = ParagraphStyle(
        "ReportHeading",
        parent=styles["Heading2"],
        fontSize=15,
        leading=19,
        spaceBefore=14,
        spaceAfter=8,
    )

    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontSize=9.5,
        leading=14,
        spaceAfter=6,
    )

    story = []

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "InsightAI Analytics Report",
            title_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>Dataset:</b> {_safe_text(dataset_name)}",
            body_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>Generated:</b> "
            f"{datetime.now().strftime('%Y-%m-%d %H:%M')}",
            body_style,
        )
    )

    story.append(Spacer(1, 12))

    # --------------------------------------------------------
    # KPI TABLE
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Executive Overview",
            heading_style,
        )
    )

    kpi_data = [
        ["Metric", "Value"],
        ["Rows", f"{report_data['rows']:,}"],
        ["Columns", f"{report_data['columns']:,}"],
        [
            "Completeness",
            f"{report_data['completeness']:.1f}%",
        ],
        [
            "Missing Cells",
            f"{report_data['missing_cells']:,}",
        ],
        [
            "Duplicate Rows",
            f"{report_data['duplicate_rows']:,}",
        ],
    ]

    story.append(
        _pdf_table(
            kpi_data,
            widths=[3.2 * inch, 2.5 * inch],
        )
    )

    story.append(Spacer(1, 12))

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Executive Summary",
            heading_style,
        )
    )

    for item in generate_executive_summary(
        df,
        report_data,
    ):

        story.append(
            Paragraph(
                f"• {_safe_text(item)}",
                body_style,
            )
        )

    # --------------------------------------------------------
    # FINDINGS
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Key Findings",
            heading_style,
        )
    )

    for finding in generate_report_findings(df):

        story.append(
            Paragraph(
                f"• {_safe_text(finding)}",
                body_style,
            )
        )

    # --------------------------------------------------------
    # DATA STRUCTURE
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Data Structure",
            heading_style,
        )
    )

    structure_data = [
        [
            "Column",
            "Data Type",
            "Missing",
            "Unique",
        ]
    ]

    for column in df.columns:

        structure_data.append(
            [
                str(column),
                str(df[column].dtype),
                str(int(df[column].isna().sum())),
                str(int(df[column].nunique(dropna=True))),
            ]
        )

    story.append(
        _pdf_table(
            structure_data,
            widths=[
                2.0 * inch,
                1.4 * inch,
                1.0 * inch,
                1.0 * inch,
            ],
        )
    )

    story.append(PageBreak())

    # --------------------------------------------------------
    # NUMERIC STATISTICS
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Numeric Statistics",
            heading_style,
        )
    )

    stats = numeric_statistics(df)

    if not stats.empty:

        columns = [
            "Column",
            "count",
            "mean",
            "std",
            "min",
            "max",
        ]

        columns = [
            col for col in columns
            if col in stats.columns
        ]

        stats_data = [
            columns
        ]

        for _, row in stats[columns].head(30).iterrows():

            stats_data.append(
                [
                    _safe_text(row[col])
                    for col in columns
                ]
            )

        story.append(
            _pdf_table(
                stats_data
            )
        )

    else:

        story.append(
            Paragraph(
                "No numeric columns were detected.",
                body_style,
            )
        )

    # --------------------------------------------------------
    # CATEGORICAL ANALYSIS
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Categorical Analysis",
            heading_style,
        )
    )

    categories = categorical_statistics(df)

    if not categories.empty:

        category_data = [
            list(categories.columns)
        ]

        for _, row in categories.head(30).iterrows():

            category_data.append(
                [
                    _safe_text(row[col])
                    for col in categories.columns
                ]
            )

        story.append(
            _pdf_table(
                category_data
            )
        )

    else:

        story.append(
            Paragraph(
                "No categorical columns were detected.",
                body_style,
            )
        )

    # --------------------------------------------------------
    # CORRELATIONS
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "Strongest Correlations",
            heading_style,
        )
    )

    correlations = correlation_summary(df)

    if not correlations.empty:

        correlation_data = [
            list(correlations.columns)
        ]

        for _, row in correlations.head(20).iterrows():

            correlation_data.append(
                [
                    _safe_text(row[col])
                    for col in correlations.columns
                ]
            )

        story.append(
            _pdf_table(
                correlation_data
            )
        )

    else:

        story.append(
            Paragraph(
                "Insufficient numeric fields for correlation analysis.",
                body_style,
            )
        )

    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    story.append(Spacer(1, 18))

    story.append(
        Paragraph(
            "Generated by InsightAI — AI-Powered Data Analytics & Decision Intelligence Platform",
            body_style,
        )
    )

    document.build(story)

    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# DOCX REPORT
# ============================================================

def generate_docx_report(
    df,
    dataset_name="InsightAI Dataset",
):

    report_data = build_report_data(df)

    document = Document()

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    title = document.add_heading(
        "InsightAI Analytics Report",
        level=0,
    )

    title.alignment = 1

    document.add_paragraph(
        f"Dataset: {dataset_name}"
    )

    document.add_paragraph(
        f"Generated: "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )

    # --------------------------------------------------------
    # EXECUTIVE OVERVIEW
    # --------------------------------------------------------

    document.add_heading(
        "Executive Overview",
        level=1,
    )

    table = document.add_table(
        rows=1,
        cols=2,
    )

    table.style = "Table Grid"

    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"

    kpis = [
        (
            "Rows",
            f"{report_data['rows']:,}",
        ),
        (
            "Columns",
            f"{report_data['columns']:,}",
        ),
        (
            "Completeness",
            f"{report_data['completeness']:.1f}%",
        ),
        (
            "Missing Cells",
            f"{report_data['missing_cells']:,}",
        ),
        (
            "Duplicate Rows",
            f"{report_data['duplicate_rows']:,}",
        ),
    ]

    for metric, value in kpis:

        cells = table.add_row().cells

        cells[0].text = metric
        cells[1].text = value

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    document.add_heading(
        "Executive Summary",
        level=1,
    )

    for item in generate_executive_summary(
        df,
        report_data,
    ):

        document.add_paragraph(
            item,
            style="List Bullet",
        )

    # --------------------------------------------------------
    # FINDINGS
    # --------------------------------------------------------

    document.add_heading(
        "Key Findings",
        level=1,
    )

    for finding in generate_report_findings(df):

        document.add_paragraph(
            finding,
            style="List Bullet",
        )

    # --------------------------------------------------------
    # DATA STRUCTURE
    # --------------------------------------------------------

    document.add_heading(
        "Data Structure",
        level=1,
    )

    structure_table = document.add_table(
        rows=1,
        cols=4,
    )

    structure_table.style = "Table Grid"

    headers = [
        "Column",
        "Data Type",
        "Missing",
        "Unique",
    ]

    for i, header in enumerate(headers):

        structure_table.rows[0].cells[i].text = header

    for column in df.columns:

        cells = structure_table.add_row().cells

        cells[0].text = str(column)
        cells[1].text = str(df[column].dtype)
        cells[2].text = str(
            int(df[column].isna().sum())
        )
        cells[3].text = str(
            int(df[column].nunique(dropna=True))
        )

    # --------------------------------------------------------
    # NUMERIC STATISTICS
    # --------------------------------------------------------

    document.add_heading(
        "Numeric Statistics",
        level=1,
    )

    stats = numeric_statistics(df)

    if not stats.empty:

        selected_columns = [
            "Column",
            "count",
            "mean",
            "std",
            "min",
            "max",
        ]

        selected_columns = [
            col for col in selected_columns
            if col in stats.columns
        ]

        stats_table = document.add_table(
            rows=1,
            cols=len(selected_columns),
        )

        stats_table.style = "Table Grid"

        for i, col in enumerate(selected_columns):

            stats_table.rows[0].cells[i].text = col

        for _, row in stats[selected_columns].head(30).iterrows():

            cells = stats_table.add_row().cells

            for i, col in enumerate(selected_columns):

                cells[i].text = _safe_text(
                    row[col]
                )

    else:

        document.add_paragraph(
            "No numeric columns were detected."
        )

    # --------------------------------------------------------
    # CATEGORICAL ANALYSIS
    # --------------------------------------------------------

    document.add_heading(
        "Categorical Analysis",
        level=1,
    )

    categories = categorical_statistics(df)

    if not categories.empty:

        category_table = document.add_table(
            rows=1,
            cols=len(categories.columns),
        )

        category_table.style = "Table Grid"

        for i, col in enumerate(categories.columns):

            category_table.rows[0].cells[i].text = col

        for _, row in categories.head(30).iterrows():

            cells = category_table.add_row().cells

            for i, col in enumerate(categories.columns):

                cells[i].text = _safe_text(
                    row[col]
                )

    else:

        document.add_paragraph(
            "No categorical columns were detected."
        )

    # --------------------------------------------------------
    # CORRELATIONS
    # --------------------------------------------------------

    document.add_heading(
        "Strongest Correlations",
        level=1,
    )

    correlations = correlation_summary(df)

    if not correlations.empty:

        corr_table = document.add_table(
            rows=1,
            cols=3,
        )

        corr_table.style = "Table Grid"

        headers = [
            "Field 1",
            "Field 2",
            "Correlation",
        ]

        for i, header in enumerate(headers):

            corr_table.rows[0].cells[i].text = header

        for _, row in correlations.head(20).iterrows():

            cells = corr_table.add_row().cells

            cells[0].text = _safe_text(
                row["Field 1"]
            )

            cells[1].text = _safe_text(
                row["Field 2"]
            )

            cells[2].text = _safe_text(
                row["Correlation"]
            )

    else:

        document.add_paragraph(
            "Insufficient numeric fields for correlation analysis."
        )

    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    document.add_paragraph("")

    footer = document.add_paragraph()

    footer.add_run(
        "Generated by InsightAI — AI-Powered Data Analytics "
        "& Decision Intelligence Platform"
    ).bold = True

    buffer = io.BytesIO()

    document.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()
